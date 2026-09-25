"""Incremental OneBSS report downloader used by the Employ dashboard.

The native OneBSS export for a long date range regularly exceeds the gateway
timeout.  This module requests the report one day at a time, merges every
successful workbook immediately, and keeps a small checkpoint so a stopped
run can continue without downloading completed days again.
"""

from __future__ import annotations

import base64
import hashlib
import io
import json
import os
import re
import threading
import time
import unicodedata
import uuid
from datetime import date, datetime, timedelta
from typing import Any, Dict, Iterable, List, Optional, Sequence, Tuple

import requests
import urllib3
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


urllib3.disable_warnings(urllib3.exceptions.InsecureRequestWarning)

REPORT_ID = 47363
REPORT_PATH = "KHAC/KHAC/RP_BC_BSS_120555"
REPORT_TITLE = "Báo cáo đơn hàng di động (Mô hình SX mới)"
API_BASE = "https://api-onebss.vnpt.vn"

# The values below are the exact options displayed in the confirmed OneBSS
# screen: HCM / VNPT Employee / Mobile / Prepaid / All payment methods.
FIXED_PARAMS = {
    "P_PHANVUNG_ID": "28",
    "P_UNGDUNG": "19",
    "P_DICHVUVT_ID": "2",
    "P_LOAITB_ID": "21",
    "P_SDT_KHMUA": "",
    "P_SDT_NV_XLYHD": "",
    "P_HINHTHUC_TT": "0",
}

_KNOWN_HEADER_KEYS = {
    "DONVI",
    "PHONGBANHANG",
    "NGUONHANG",
    "MA_DH_CHAM",
    "LOAIDON",
    "HINHTHUC_DH",
    "MADONHANG",
    "MADONHANG_TRANGTHAI",
    "MA_ID_ONEBSS",
    "MA_ID_ONEBSS_TRANGTHAI",
    "NHOMSANPHAM",
    "SANPHAM",
}


class OneBssReportError(RuntimeError):
    """A recoverable report/session/export error."""


def _normalise_header(value: Any) -> str:
    text = unicodedata.normalize("NFD", str(value or ""))
    text = "".join(ch for ch in text if unicodedata.category(ch) != "Mn")
    return re.sub(r"[^A-Z0-9]+", "_", text.upper()).strip("_")


def _decode_jwt_payload(token: str) -> Dict[str, Any]:
    try:
        payload = token.split(".")[1]
        payload += "=" * (-len(payload) % 4)
        decoded = base64.urlsafe_b64decode(payload.encode("ascii"))
        data = json.loads(decoded.decode("utf-8"))
        return data if isinstance(data, dict) else {}
    except (ValueError, IndexError, TypeError, json.JSONDecodeError):
        return {}


def load_onebss_session(session_file: str) -> Tuple[str, Dict[str, Any]]:
    """Read the latest saved OneBSS token without ever exposing it to the UI."""
    try:
        with open(session_file, "r", encoding="utf-8-sig") as handle:
            root = json.load(handle)
    except FileNotFoundError as exc:
        raise OneBssReportError(
            f"Không tìm thấy phiên OneBSS: {session_file}"
        ) from exc
    except (OSError, json.JSONDecodeError) as exc:
        raise OneBssReportError(f"Không đọc được phiên OneBSS: {exc}") from exc

    active = root.get("sessions", {}).get("active", {}) if isinstance(root, dict) else {}
    token_value = active.get("token") if isinstance(active, dict) else None
    if isinstance(token_value, dict):
        token = str(token_value.get("access_token") or "").strip()
    else:
        token = str(token_value or active.get("access_token") or "").strip()
    if not token:
        raise OneBssReportError("File phiên chưa có access_token; hãy đăng nhập lại OneBSS.")

    claims = _decode_jwt_payload(token)
    expires_at = int(claims.get("exp") or 0)
    if expires_at and expires_at <= int(time.time()) + 15:
        raise OneBssReportError(
            "Phiên OneBSS đã hết hạn; hãy đăng nhập lại rồi bấm Tiếp tục."
        )
    return token, claims


def validate_access_token(token: str) -> Tuple[str, Dict[str, Any]]:
    token = str(token or "").strip()
    if token.lower().startswith("bearer "):
        token = token[7:].strip()
    if not token:
        raise OneBssReportError("Tài khoản được chọn chưa có access token.")
    claims = _decode_jwt_payload(token)
    expires_at = int(claims.get("exp") or 0)
    if expires_at and expires_at <= int(time.time()) + 15:
        raise OneBssReportError(
            "Token tài khoản được chọn đã hết hạn; hãy dán token mới rồi Tiếp tục."
        )
    return token, claims


def _iter_days(start_day: date, end_day: date) -> Iterable[date]:
    cursor = start_day
    while cursor <= end_day:
        yield cursor
        cursor += timedelta(days=1)


def _is_xlsx_response(response: requests.Response) -> bool:
    content_type = str(response.headers.get("Content-Type") or "").lower()
    return (
        response.content.startswith(b"PK")
        or "spreadsheetml" in content_type
        or "application/vnd.ms-excel" in content_type
    )


def _response_error(response: requests.Response) -> str:
    try:
        payload = response.json()
        if isinstance(payload, dict):
            return str(
                payload.get("message_detail")
                or payload.get("message")
                or payload.get("error_code")
                or response.status_code
            )
    except ValueError:
        pass
    text = (response.text or "").strip()
    return text[:500] if text else f"HTTP {response.status_code}"


def download_report_day(
    session_file: str,
    report_day: date,
    *,
    access_token: str = "",
    timeout_seconds: int = 180,
    http_session: Optional[requests.Session] = None,
) -> bytes:
    """Download one calendar day using the same direct-export API as /report/bi."""
    token, claims = (
        validate_access_token(access_token)
        if access_token else load_onebss_session(session_file)
    )
    params = dict(FIXED_PARAMS)
    rendered_day = report_day.strftime("%d/%m/%Y")
    params.update({
        "P_TUNGAY": rendered_day,
        "P_DENNGAY": rendered_day,
        "username": str(claims.get("user_name") or claims.get("preferred_username") or ""),
    })
    body = {"baocao_id": REPORT_ID, "params": params}
    headers = {
        "Authorization": f"Bearer {token}",
        "apiKey": "x",
        "Content-Type": "application/json",
    }
    client = http_session or requests.Session()
    request_id = f"{int(time.time() * 1000)}{uuid.uuid4().hex[:6]}"
    url = f"{API_BASE}/web-report/report/bi/run_v5?requestId={request_id}"

    try:
        response = client.post(
            url,
            json=body,
            headers=headers,
            verify=False,
            timeout=(20, max(30, int(timeout_seconds))),
        )
    except requests.Timeout as exc:
        raise OneBssReportError(
            f"OneBSS quá thời gian khi xuất ngày {rendered_day}"
        ) from exc
    except requests.RequestException as exc:
        raise OneBssReportError(f"Lỗi kết nối OneBSS: {exc}") from exc

    # Some report jobs are queued.  The first response supplies either a
    # Location header or a short relative path in its body.
    if response.status_code == 202:
        location = str(response.headers.get("Location") or "").strip()
        if not location:
            location = response.text.strip().strip('"')
        if not location:
            raise OneBssReportError("OneBSS đã tạo tác vụ nhưng không trả đường dẫn kết quả.")
        if location.startswith("/"):
            location = f"{API_BASE}/web-report{location}"
        elif not location.startswith("http"):
            location = f"{API_BASE}/web-report/{location.lstrip('/')}"
        deadline = time.monotonic() + max(60, int(timeout_seconds))
        while time.monotonic() < deadline:
            time.sleep(10)
            try:
                response = client.post(
                    location,
                    json=body,
                    headers=headers,
                    verify=False,
                    timeout=(20, 60),
                )
            except requests.Timeout:
                continue
            if response.status_code != 202:
                break

    if response.status_code in (401, 403):
        raise OneBssReportError("Phiên OneBSS không còn hợp lệ; hãy đăng nhập lại.")
    if response.status_code >= 400:
        raise OneBssReportError(_response_error(response))
    if not _is_xlsx_response(response):
        raise OneBssReportError(
            "OneBSS không trả file Excel: " + _response_error(response)
        )
    return response.content


def fetch_report_day_table(
    session_file: str,
    report_day: date,
    *,
    access_token: str = "",
    timeout_seconds: int = 90,
    http_session: Optional[requests.Session] = None,
) -> Tuple[List[str], List[List[Any]]]:
    """Read one day through the fast grid API and return a local table.

    OneBSS's run_v5 endpoint renders a complete XLSX on the server and can
    time out even for a single busy day.  The /report/bi screen itself uses
    run_v7 for "Xem lưới"; consuming that JSON and building the final XLSX
    locally avoids the expensive server-side workbook rendering.
    """
    token, claims = (
        validate_access_token(access_token)
        if access_token else load_onebss_session(session_file)
    )
    rendered_day = report_day.strftime("%d/%m/%Y")
    params = dict(FIXED_PARAMS)
    params.update({
        "P_TUNGAY": rendered_day,
        "P_DENNGAY": rendered_day,
        "username": str(claims.get("user_name") or claims.get("preferred_username") or ""),
    })
    body = {"baocao_id": REPORT_ID, "params": params}
    headers = {
        "Authorization": f"Bearer {token}",
        "apiKey": "x",
        "Content-Type": "application/json",
    }
    client = http_session or requests.Session()
    request_id = f"{int(time.time() * 1000)}{uuid.uuid4().hex[:6]}"
    url = f"{API_BASE}/web-report/report/bi/run_v7?requestId={request_id}"
    try:
        response = client.post(
            url,
            json=body,
            headers=headers,
            verify=False,
            timeout=(20, max(30, int(timeout_seconds))),
        )
    except requests.Timeout as exc:
        raise OneBssReportError(
            f"OneBSS quá thời gian khi đọc ngày {rendered_day}"
        ) from exc
    except requests.RequestException as exc:
        raise OneBssReportError(f"Lỗi kết nối OneBSS: {exc}") from exc

    if response.status_code == 202:
        location = str(response.headers.get("Location") or "").strip()
        if not location:
            location = response.text.strip().strip('"')
        if not location:
            raise OneBssReportError("OneBSS không trả đường dẫn kết quả lưới.")
        if location.startswith("/"):
            location = f"{API_BASE}/web-report{location}"
        elif not location.startswith("http"):
            location = f"{API_BASE}/web-report/{location.lstrip('/')}"
        deadline = time.monotonic() + max(60, int(timeout_seconds))
        while time.monotonic() < deadline:
            time.sleep(5)
            response = client.post(
                location,
                json=body,
                headers=headers,
                verify=False,
                timeout=(20, 45),
            )
            if response.status_code != 202:
                break

    if response.status_code in (401, 403):
        raise OneBssReportError("Phiên OneBSS không còn hợp lệ; hãy đăng nhập lại.")
    if response.status_code >= 400:
        raise OneBssReportError(_response_error(response))
    try:
        payload = response.json()
    except ValueError as exc:
        raise OneBssReportError("OneBSS trả dữ liệu lưới không phải JSON.") from exc
    if not isinstance(payload, dict):
        raise OneBssReportError("Dữ liệu lưới OneBSS không đúng cấu trúc.")
    code = str(payload.get("error_code") or "")
    if code and code not in ("0", "200", "BSS-00000000"):
        raise OneBssReportError(
            str(payload.get("message_detail") or payload.get("message") or code)
        )
    records = payload.get("data") or []
    if isinstance(records, str):
        try:
            records = json.loads(records)
        except json.JSONDecodeError:
            records = []
    records = [item for item in records if isinstance(item, dict)] if isinstance(records, list) else []
    if not records or not records[0]:
        return [], []

    columns = list(records[0].keys())
    type_names = {
        "STRING", "NUMBER", "INTEGER", "DECIMAL", "DATE", "DATETIME",
        "TIMESTAMP", "BOOLEAN", "DOUBLE", "FLOAT", "LONG", "SHORT",
    }
    first_values = [str(records[0].get(column) or "").strip().upper() for column in columns]
    typed = sum(value in type_names for value in first_values)
    start_index = 1 if first_values and typed >= max(1, len(first_values) // 2) else 0
    rows = [
        [record.get(column, "") for column in columns]
        for record in records[start_index:]
        if any(value not in (None, "") for value in record.values())
    ]
    return columns, rows


def extract_report_table(workbook_bytes: bytes) -> Tuple[List[str], List[List[Any]]]:
    """Extract the technical header row and data rows from a report workbook."""
    try:
        workbook = load_workbook(io.BytesIO(workbook_bytes), read_only=True, data_only=True)
    except Exception as exc:
        raise OneBssReportError(f"File Excel OneBSS không hợp lệ: {exc}") from exc

    try:
        sheets = [workbook[name] for name in workbook.sheetnames]
        sheet = max(sheets, key=lambda item: item.max_row * max(1, item.max_column))
        preview = list(
            sheet.iter_rows(
                min_row=1,
                max_row=min(max(1, sheet.max_row), 40),
                values_only=True,
            )
        )
        header_index = None
        best_score = 0
        for index, row in enumerate(preview):
            keys = {_normalise_header(value) for value in row if value not in (None, "")}
            score = len(keys & _KNOWN_HEADER_KEYS)
            if score > best_score:
                best_score = score
                header_index = index

        # Empty exports can contain a report title and parameters but no grid
        # header.  Do not mistake those descriptive rows for business data.
        if header_index is None or best_score < 2:
            return [], []

        raw_header = list(preview[header_index])
        last_column = max(
            (index for index, value in enumerate(raw_header) if value not in (None, "")),
            default=-1,
        )
        if last_column < 0:
            return [], []
        raw_header = raw_header[: last_column + 1]
        headers: List[str] = []
        used: Dict[str, int] = {}
        for column, value in enumerate(raw_header, start=1):
            base = str(value or "").strip() or f"CỘT_{column}"
            count = used.get(base, 0) + 1
            used[base] = count
            headers.append(base if count == 1 else f"{base}_{count}")

        rows: List[List[Any]] = []
        header_keys = [_normalise_header(value) for value in headers]
        for row in sheet.iter_rows(min_row=header_index + 2, values_only=True):
            values = list(row[: len(headers)])
            if not any(value not in (None, "") for value in values):
                continue
            if [_normalise_header(value) for value in values] == header_keys:
                continue
            rows.append(values)
        return headers, rows
    finally:
        workbook.close()


def _row_digest(values: Sequence[Any]) -> str:
    serialised = json.dumps(
        ["" if value is None else str(value) for value in values],
        ensure_ascii=False,
        separators=(",", ":"),
    )
    return hashlib.sha256(serialised.encode("utf-8")).hexdigest()


def _atomic_save_workbook(workbook: Workbook, path: str) -> None:
    os.makedirs(os.path.dirname(path), exist_ok=True)
    temporary = f"{path}.{uuid.uuid4().hex}.tmp.xlsx"
    try:
        workbook.save(temporary)
        os.replace(temporary, path)
    finally:
        if os.path.exists(temporary):
            try:
                os.remove(temporary)
            except OSError:
                pass


def _atomic_save_json(payload: Dict[str, Any], path: str) -> None:
    os.makedirs(os.path.dirname(path), exist_ok=True)
    temporary = f"{path}.{uuid.uuid4().hex}.tmp"
    try:
        with open(temporary, "w", encoding="utf-8") as handle:
            json.dump(payload, handle, ensure_ascii=False, indent=2)
        os.replace(temporary, path)
    finally:
        if os.path.exists(temporary):
            try:
                os.remove(temporary)
            except OSError:
                pass


def merge_report_rows(output_path: str, headers: Sequence[str], rows: Sequence[Sequence[Any]]) -> int:
    """Merge one day's data into the single user-visible output workbook."""
    if os.path.isfile(output_path):
        workbook = load_workbook(output_path)
        sheet = workbook["Du_lieu"] if "Du_lieu" in workbook.sheetnames else workbook.active
    else:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Du_lieu"

    try:
        existing_headers = [str(cell.value or "").strip() for cell in sheet[1]] if sheet.max_row else []
        if not any(existing_headers):
            if sheet.max_row:
                sheet.delete_rows(1, sheet.max_row)
            existing_headers = []
        elif existing_headers == ["Thông báo"] and headers:
            sheet.delete_rows(1, sheet.max_row)
            existing_headers = []
        final_headers = list(existing_headers)
        for header in headers:
            if header not in final_headers:
                final_headers.append(header)

        if final_headers and not existing_headers:
            sheet.append(final_headers)
        elif len(final_headers) > len(existing_headers):
            for index, header in enumerate(final_headers, start=1):
                sheet.cell(1, index, header)

        existing_digests = set()
        if sheet.max_row > 1 and final_headers:
            for current in sheet.iter_rows(min_row=2, max_col=len(final_headers), values_only=True):
                existing_digests.add(_row_digest(current))

        incoming_index = {name: index for index, name in enumerate(headers)}
        appended = 0
        for row in rows:
            mapped = [
                row[incoming_index[name]] if name in incoming_index and incoming_index[name] < len(row) else ""
                for name in final_headers
            ]
            digest = _row_digest(mapped)
            if digest in existing_digests:
                continue
            sheet.append(mapped)
            existing_digests.add(digest)
            appended += 1

        if not final_headers:
            sheet.cell(1, 1, "Thông báo")
            sheet.cell(2, 1, "Không có dữ liệu trong phạm vi đã chọn")
            final_headers = ["Thông báo"]

        header_fill = PatternFill("solid", fgColor="1F4E78")
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=False)
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = f"A1:{get_column_letter(len(final_headers))}{max(1, sheet.max_row)}"
        sheet.sheet_view.showGridLines = False

        for column_index, header in enumerate(final_headers, start=1):
            width = min(52, max(12, len(str(header)) + 3))
            for cells in sheet.iter_rows(
                min_row=2,
                max_row=min(sheet.max_row, 400),
                min_col=column_index,
                max_col=column_index,
            ):
                value = cells[0].value
                if value not in (None, ""):
                    width = min(52, max(width, len(str(value)) + 2))
                cells[0].alignment = Alignment(vertical="center", wrap_text=False)
            sheet.column_dimensions[get_column_letter(column_index)].width = width
        sheet.row_dimensions[1].height = 24
        _atomic_save_workbook(workbook, output_path)
        return appended
    finally:
        workbook.close()


class OneBssReportBatchManager:
    """Own one background report run and expose a small serialisable status."""

    def __init__(self, session_file: str, output_dir: str):
        self.session_file = os.path.abspath(session_file)
        self.output_dir = os.path.abspath(output_dir)
        self._lock = threading.RLock()
        self._stop_event = threading.Event()
        self._thread: Optional[threading.Thread] = None
        self._run_token = ""
        self._state: Dict[str, Any] = {
            "running": False,
            "status": "idle",
            "message": "Chưa chạy",
            "completed_days": 0,
            "total_days": 0,
            "rows": 0,
            "output_path": "",
            "current_day": "",
        }

    def status(self) -> Dict[str, Any]:
        with self._lock:
            if not self._state.get("running") and not self._state.get("output_path"):
                try:
                    candidates = [
                        entry.path for entry in os.scandir(self.output_dir)
                        if entry.is_file()
                        and entry.name.startswith("Bao_Cao_OneBSS_VNPT_Employee_")
                        and entry.name.lower().endswith(".xlsx")
                    ]
                except FileNotFoundError:
                    candidates = []
                if candidates:
                    latest = max(candidates, key=os.path.getmtime)
                    self._state.update({
                        "status": "available",
                        "message": "File hoàn tất gần nhất",
                        "output_path": latest,
                    })
            state = dict(self._state)
        state["session_file"] = self.session_file
        state["output_dir"] = self.output_dir
        state["download_ready"] = bool(
            state.get("output_path") and os.path.isfile(str(state["output_path"]))
        )
        return state

    def start(
        self,
        start_day: date,
        end_day: date,
        *,
        access_token: str = "",
        account_username: str = "",
    ) -> Dict[str, Any]:
        if start_day > end_day:
            raise OneBssReportError("Từ ngày phải nhỏ hơn hoặc bằng Đến ngày.")
        with self._lock:
            if self._thread and self._thread.is_alive():
                raise OneBssReportError("Báo cáo đang chạy.")
            self._stop_event.clear()
            self._run_token = str(access_token or "").strip()
            self._state = {
                "running": True,
                "status": "running",
                "message": "Đang chuẩn bị",
                "start_date": start_day.isoformat(),
                "end_date": end_day.isoformat(),
                "completed_days": 0,
                "total_days": (end_day - start_day).days + 1,
                "rows": 0,
                "output_path": "",
                "current_day": "",
                "account_username": str(account_username or "onebss_session.json"),
            }
            self._thread = threading.Thread(
                target=self._run,
                args=(start_day, end_day, str(account_username or "onebss_session")),
                name="onebss-report-batch",
                daemon=True,
            )
            self._thread.start()
        return self.status()

    def stop(self) -> Dict[str, Any]:
        self._stop_event.set()
        with self._lock:
            if self._state.get("running"):
                self._state["message"] = "Đang dừng sau ngày hiện tại…"
        return self.status()

    def _set_state(self, **changes: Any) -> None:
        with self._lock:
            self._state.update(changes)

    @staticmethod
    def _account_slug(account_username: str) -> str:
        slug = re.sub(r"[^A-Za-z0-9_.-]+", "_", str(account_username or "").strip())
        return slug[:50] or "onebss_session"

    def _checkpoint_path(self, start_day: date, end_day: date, account_username: str) -> str:
        slug = self._account_slug(account_username)
        name = f".onebss_employee_{slug}_{start_day:%Y%m%d}_{end_day:%Y%m%d}.checkpoint.json"
        return os.path.join(self.output_dir, name)

    def _new_output_path(self, start_day: date, end_day: date, account_username: str) -> str:
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        slug = self._account_slug(account_username)
        name = (
            f"Bao_Cao_OneBSS_VNPT_Employee_{slug}_{start_day:%Y%m%d}_{end_day:%Y%m%d}_{timestamp}.xlsx"
        )
        return os.path.join(self.output_dir, name)

    def _load_checkpoint(self, checkpoint_path: str) -> Dict[str, Any]:
        try:
            with open(checkpoint_path, "r", encoding="utf-8") as handle:
                payload = json.load(handle)
            return payload if isinstance(payload, dict) else {}
        except (FileNotFoundError, OSError, json.JSONDecodeError):
            return {}

    def _run(self, start_day: date, end_day: date, account_username: str) -> None:
        os.makedirs(self.output_dir, exist_ok=True)
        checkpoint_path = self._checkpoint_path(start_day, end_day, account_username)
        checkpoint = self._load_checkpoint(checkpoint_path)
        output_path = str(checkpoint.get("output_path") or "")
        try:
            output_owned = os.path.commonpath([
                os.path.abspath(output_path), self.output_dir
            ]) == self.output_dir
        except (ValueError, TypeError):
            output_owned = False
        if not output_path or not output_owned:
            output_path = self._new_output_path(start_day, end_day, account_username)
        completed = {
            str(value) for value in checkpoint.get("completed_days", []) if value
        }
        rows_total = int(checkpoint.get("rows", 0) or 0)
        if completed and rows_total > 0 and not os.path.isfile(output_path):
            # A workbook may have been manually moved/deleted.  Never skip its
            # completed days unless the actual accumulated data still exists.
            completed.clear()
            rows_total = 0
        all_days = list(_iter_days(start_day, end_day))
        self._set_state(
            output_path=output_path,
            completed_days=len(completed),
            rows=rows_total,
        )

        try:
            # Fail quickly before starting the first slow report request.
            if self._run_token:
                validate_access_token(self._run_token)
            else:
                load_onebss_session(self.session_file)
            for report_day in all_days:
                day_key = report_day.isoformat()
                if day_key in completed:
                    continue
                if self._stop_event.is_set():
                    self._set_state(
                        running=False,
                        status="stopped",
                        message="Đã dừng; lần chạy sau sẽ tiếp tục từ checkpoint.",
                    )
                    self._run_token = ""
                    return

                self._set_state(
                    current_day=day_key,
                    message=f"Đang tải {report_day:%d/%m/%Y}",
                )
                last_error: Optional[Exception] = None
                day_table: Optional[Tuple[List[str], List[List[Any]]]] = None
                for attempt, delay in enumerate((0, 8, 20), start=1):
                    if delay and self._stop_event.wait(delay):
                        break
                    try:
                        day_table = fetch_report_day_table(
                            self.session_file,
                            report_day,
                            access_token=self._run_token,
                            timeout_seconds=90,
                        )
                        last_error = None
                        break
                    except OneBssReportError as exc:
                        last_error = exc
                        self._set_state(
                            message=(
                                f"Ngày {report_day:%d/%m/%Y} lỗi lần {attempt}/3: {exc}"
                            )
                        )
                if day_table is None:
                    if self._stop_event.is_set():
                        self._set_state(
                            running=False,
                            status="stopped",
                            message="Đã dừng; lần chạy sau sẽ tiếp tục từ checkpoint.",
                        )
                        self._run_token = ""
                        return
                    raise last_error or OneBssReportError("Không tải được dữ liệu.")

                headers, rows = day_table
                if headers:
                    rows_total += merge_report_rows(output_path, headers, rows)
                completed.add(day_key)
                checkpoint = {
                    "report_id": REPORT_ID,
                    "report_path": REPORT_PATH,
                    "start_date": start_day.isoformat(),
                    "end_date": end_day.isoformat(),
                    "output_path": output_path,
                    "completed_days": sorted(completed),
                    "rows": rows_total,
                    "updated_at": datetime.now().isoformat(timespec="seconds"),
                }
                _atomic_save_json(checkpoint, checkpoint_path)
                self._set_state(
                    completed_days=len(completed),
                    rows=rows_total,
                    message=(
                        f"Đã xong {len(completed)}/{len(all_days)} ngày · {rows_total} dòng"
                    ),
                )

            if not os.path.isfile(output_path):
                merge_report_rows(output_path, [], [])
            try:
                os.remove(checkpoint_path)
            except FileNotFoundError:
                pass
            self._set_state(
                running=False,
                status="completed",
                current_day="",
                message=f"Hoàn tất {len(all_days)} ngày · {rows_total} dòng",
                completed_days=len(all_days),
                rows=rows_total,
            )
            self._run_token = ""
        except Exception as exc:
            self._set_state(
                running=False,
                status="error",
                message=str(exc),
            )
            self._run_token = ""

