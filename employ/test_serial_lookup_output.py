import os
import tempfile
import unittest
from unittest.mock import patch

from flask import session
from openpyxl import load_workbook

import app as app_module


class SerialLookupOutputTests(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory(prefix='vnpt_serial_lookup_test_')
        self.batch_dir = os.path.join(self.temp_dir.name, 'Tra_Cuu_Seri')
        self.input_file = os.path.join(self.batch_dir, 'Tra_Cuu_Seri_Input.xlsx')
        self.output_file = os.path.join(self.batch_dir, 'Tra_Cuu_Seri_Output.xlsx')
        self.patchers = [
            patch.object(app_module, 'SERIAL_LOOKUP_BATCH_DIR', self.batch_dir),
            patch.object(app_module, 'SERIAL_LOOKUP_INPUT_FILE', self.input_file),
            patch.object(app_module, 'SERIAL_LOOKUP_OUTPUT_FILE', self.output_file),
        ]
        for patcher in self.patchers:
            patcher.start()

    def tearDown(self):
        for patcher in reversed(self.patchers):
            patcher.stop()
        self.temp_dir.cleanup()

    def _append(self, serial, phone, entry_id):
        payload = {
            'table': [list(app_module.SERIAL_LOOKUP_OUTPUT_HEADERS), [serial, phone]],
            'entry_ids': [entry_id],
        }
        with app_module.app.test_request_context(json=payload):
            session['username'] = 'tester'
            response = app_module.api_serial_lookup_append_output.__wrapped__()
        return response.get_json()

    def test_files_use_one_input_column_and_two_output_columns(self):
        app_module._ensure_serial_lookup_files()

        input_book = load_workbook(self.input_file, read_only=True, data_only=True)
        input_sheet = input_book['Input']
        self.assertEqual([cell.value for cell in input_sheet[1]], ['MSIN / Seri SIM'])
        input_book.close()

        output_book = load_workbook(self.output_file, read_only=True, data_only=True)
        output_sheet = output_book['Tra_Cuu_Seri']
        self.assertEqual(
            [cell.value for cell in output_sheet[1]][:2],
            ['MSIN / Seri SIM', 'SĐT'],
        )
        output_book.close()

    def test_each_finished_row_is_saved_and_retry_updates_same_row(self):
        first = self._append('1170348501', '', 'run-a-row-1')
        retry = self._append('1170348501', '0919138835', 'run-a-row-1')
        duplicate_retry = self._append('1170348501', '0919138835', 'run-a-row-1')

        self.assertEqual(first['appended'], 1)
        self.assertEqual(retry['updated'], 1)
        self.assertEqual(duplicate_retry['skipped'], 1)

        workbook = load_workbook(self.output_file, data_only=False)
        sheet = workbook['Tra_Cuu_Seri']
        self.assertEqual(sheet.max_row, 2)
        self.assertEqual(sheet['A2'].value, '1170348501')
        self.assertEqual(sheet['B2'].value, '0919138835')
        self.assertEqual(sheet['A2'].data_type, 's')
        self.assertEqual(sheet['B2'].data_type, 's')
        workbook.close()


if __name__ == '__main__':
    unittest.main()
