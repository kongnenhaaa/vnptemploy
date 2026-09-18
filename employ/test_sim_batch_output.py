import os
import tempfile
import unittest
from unittest.mock import patch

from flask import session
from openpyxl import Workbook, load_workbook

import app as app_module


class SimBatchOutputTests(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory(prefix='vnpt_sim_output_test_')
        self.batch_dir = os.path.join(self.temp_dir.name, 'SIM_Kit_Batch')
        self.input_file = os.path.join(self.batch_dir, 'SIM_Kit_Input.xlsx')
        self.output_file = os.path.join(self.batch_dir, 'SIM_Kit_Output.xlsx')
        self.patchers = [
            patch.object(app_module, 'SIM_BATCH_DIR', self.batch_dir),
            patch.object(app_module, 'SIM_BATCH_INPUT_FILE', self.input_file),
            patch.object(app_module, 'SIM_BATCH_OUTPUT_FILE', self.output_file),
            patch.object(app_module, '_legacy_sim_batch_dir', os.path.join(self.temp_dir.name, 'legacy')),
        ]
        for patcher in self.patchers:
            patcher.start()

    def tearDown(self):
        for patcher in reversed(self.patchers):
            patcher.stop()
        self.temp_dir.cleanup()

    def _append(self, row, entry_id):
        payload = {
            'table': [list(app_module.SIM_BATCH_OUTPUT_HEADERS), row],
            'entry_ids': [entry_id],
        }
        with app_module.app.test_request_context(json=payload):
            session['username'] = 'tester'
            response = app_module.api_sim_batch_append_output.__wrapped__()
        return response.get_json()

    def test_each_row_appends_to_one_cumulative_file_and_retry_is_idempotent(self):
        first = [
            '84911111111', '8984000001', 'Thành công', '4911598',
            'Khách hàng tự đăng ký', 'Nguyễn Văn An', '12 Phường Cầu Kiệu',
            75000, '13/09/2026 08:00:00', 'tester',
        ]
        second = [
            '84922222222', '8984000002', 'Thành công', '4911599',
            'Nhân viên hỗ trợ', 'Trần Thị Bình', '14 Phường Cầu Kiệu',
            75000, '13/09/2026 08:01:00', 'tester',
        ]

        first_result = self._append(first, 'run-a-row-1')
        duplicate_result = self._append(first, 'run-a-row-1')
        second_result = self._append(second, 'run-b-row-1')

        self.assertEqual(first_result['appended'], 1)
        self.assertEqual(duplicate_result['appended'], 0)
        self.assertEqual(duplicate_result['skipped'], 1)
        self.assertEqual(second_result['appended'], 1)
        self.assertEqual(second_result['total_rows'], 2)
        self.assertEqual(second_result['path'], self.output_file)

        workbook = load_workbook(self.output_file, data_only=True)
        sheet = workbook['SIM_Output']
        self.assertEqual([cell.value for cell in sheet[1]], list(app_module.SIM_BATCH_OUTPUT_HEADERS))
        self.assertEqual(sheet['E1'].value, 'Loại đăng ký')
        self.assertEqual(sheet['D2'].value, '4911598')
        self.assertEqual(sheet['E2'].value, 'Khách hàng tự đăng ký')
        self.assertEqual(sheet['E3'].value, 'Nhân viên hỗ trợ')
        self.assertEqual(sheet.max_row, 3)
        self.assertEqual(workbook[app_module.SIM_BATCH_OUTPUT_META_SHEET].sheet_state, 'hidden')
        workbook.close()

    def test_legacy_output_is_migrated_without_losing_result(self):
        os.makedirs(self.batch_dir, exist_ok=True)
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = 'SIM_Output'
        sheet.append(['SĐT', 'Serial SIM', 'Kết quả', 'Thời gian', 'User chạy'])
        legacy_result = (
            'Thành công | Đơn 4911598 | Nguyễn Văn An | '
            'KH tự đăng ký ĐKTTTB | 75.000 đ | 12 Phường Cầu Kiệu'
        )
        sheet.append(['84911111111', '8984000001', legacy_result, '13/09/2026 08:00:00', 'tester'])
        workbook.save(self.output_file)
        workbook.close()

        app_module._ensure_sim_batch_files()

        migrated = load_workbook(self.output_file, data_only=True)
        sheet = migrated['SIM_Output']
        self.assertEqual(sheet['C2'].value, legacy_result)
        self.assertEqual(sheet['D2'].value, '4911598')
        self.assertEqual(sheet['E2'].value, 'Khách hàng tự đăng ký')
        self.assertEqual(sheet['F2'].value, 'Nguyễn Văn An')
        self.assertEqual(sheet['G2'].value, '12 Phường Cầu Kiệu')
        migrated.close()

    def test_completed_row_is_written_to_run_file_and_cumulative_file(self):
        with app_module.app.test_request_context(json={}):
            session['username'] = 'tester'
            started = app_module.api_sim_batch_start_output.__wrapped__().get_json()

        run_path = started['run_output_path']
        row = [
            '84911111111', '8984000001', 'Thành công', '4911598',
            'Khách hàng tự đăng ký', 'Nguyễn Văn An', '12 Phường Cầu Kiệu',
            75000, '15/09/2026 10:00:00', 'tester',
        ]
        payload = {
            'table': [list(app_module.SIM_BATCH_OUTPUT_HEADERS), row],
            'entry_ids': ['dual-output-row-1'],
            'run_output_path': run_path,
        }
        with app_module.app.test_request_context(json=payload):
            session['username'] = 'tester'
            result = app_module.api_sim_batch_append_output.__wrapped__().get_json()

        self.assertEqual(result['run_output_path'], run_path)
        self.assertEqual(result['run_total_rows'], 1)
        self.assertEqual(result['total_rows'], 1)
        for path in (run_path, self.output_file):
            workbook = load_workbook(path, data_only=True)
            self.assertEqual(workbook['SIM_Output'].max_row, 2)
            self.assertEqual(workbook['SIM_Output']['A2'].value, '84911111111')
            workbook.close()

    def test_old_per_run_files_are_merged_once_into_all_time_output(self):
        os.makedirs(self.batch_dir, exist_ok=True)
        legacy_path = os.path.join(
            self.batch_dir, 'SIM_Kit_Output_20260914_090000_abcdef.xlsx')
        legacy = Workbook()
        sheet = legacy.active
        sheet.title = 'SIM_Output'
        sheet.append(['SĐT', 'Serial SIM', 'Kết quả', 'Thời gian', 'User chạy'])
        sheet.append([
            '84944444444', '8984000044', 'Thành công | Đơn 4911600',
            '14/09/2026 09:00:00', 'old-user',
        ])
        legacy.save(legacy_path)
        legacy.close()

        app_module._ensure_sim_batch_files()
        app_module._ensure_sim_batch_files()

        total = load_workbook(self.output_file, data_only=True)
        sheet = total['SIM_Output']
        self.assertEqual(sheet.max_row, 2)
        self.assertEqual(sheet['A2'].value, '84944444444')
        self.assertEqual(sheet['B2'].value, '8984000044')
        self.assertEqual(sheet['D2'].value, '4911600')
        total.close()


if __name__ == '__main__':
    unittest.main()
