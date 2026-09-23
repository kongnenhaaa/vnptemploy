import os
import tempfile
import unittest
from unittest.mock import patch

from flask import session
from openpyxl import Workbook, load_workbook

import app as app_module


class IcocBatchOutputTests(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory(prefix='vnpt_icoc_output_test_')
        self.batch_dir = os.path.join(self.temp_dir.name, 'IC_OC_Batch')
        self.input_file = os.path.join(self.batch_dir, 'IC_OC_Input.xlsx')
        self.output_file = os.path.join(self.batch_dir, 'IC_OC_Output.xlsx')
        self.patchers = [
            patch.object(app_module, 'ICOC_BATCH_DIR', self.batch_dir),
            patch.object(app_module, 'ICOC_BATCH_INPUT_FILE', self.input_file),
            patch.object(app_module, 'ICOC_BATCH_OUTPUT_FILE', self.output_file),
        ]
        for patcher in self.patchers:
            patcher.start()

    def tearDown(self):
        for patcher in reversed(self.patchers):
            patcher.stop()
        self.temp_dir.cleanup()

    def _start_run(self, headers):
        with app_module.app.test_request_context(json={'headers': headers}):
            session['username'] = 'tester'
            response = app_module.api_icoc_batch_start_output.__wrapped__()
        return response.get_json()

    def _append(self, headers, row, entry_id, run_path):
        payload = {
            'table': [headers, row],
            'entry_ids': [entry_id],
            'run_output_path': run_path,
        }
        with app_module.app.test_request_context(json=payload):
            session['username'] = 'tester'
            response = app_module.api_icoc_batch_append_output.__wrapped__()
        return response.get_json()

    def test_each_finished_phone_is_saved_immediately_to_both_outputs(self):
        headers = ['SĐT', 'Ghi chú input', 'Kết quả', 'Thời gian', 'User chạy']
        run = self._start_run(headers)
        row = [
            '84911111111', 'ticket-01', 'Thành công | Mở IC',
            '15/09/2026 10:00:00', 'tester',
        ]

        result = self._append(headers, row, 'icoc-run-1-row-1', run['run_output_path'])

        self.assertEqual(result['run_total_rows'], 1)
        self.assertEqual(result['total_rows'], 1)
        run_book = load_workbook(run['run_output_path'], data_only=True)
        self.assertEqual(run_book['IC_OC_Output']['B2'].value, 'ticket-01')
        self.assertEqual(run_book['IC_OC_Output']['C2'].value, 'Thành công | Mở IC')
        run_book.close()
        total_book = load_workbook(self.output_file, data_only=True)
        self.assertEqual(
            [cell.value for cell in total_book['IC_OC_Output'][1]],
            list(app_module.ICOC_BATCH_OUTPUT_HEADERS),
        )
        self.assertEqual(total_book['IC_OC_Output']['A2'].value, '84911111111')
        self.assertEqual(total_book['IC_OC_Output']['B2'].value, 'Thành công | Mở IC')
        total_book.close()

    def test_retry_is_idempotent_in_run_and_cumulative_outputs(self):
        headers = ['SĐT', 'Kết quả', 'Thời gian', 'User chạy']
        run = self._start_run(headers)
        row = ['84922222222', 'Không cần đổi', '15/09/2026 10:01:00', 'tester']

        first = self._append(headers, row, 'same-entry', run['run_output_path'])
        duplicate = self._append(headers, row, 'same-entry', run['run_output_path'])

        self.assertEqual(first['appended'], 1)
        self.assertEqual(first['run_appended'], 1)
        self.assertEqual(duplicate['appended'], 0)
        self.assertEqual(duplicate['run_appended'], 0)
        self.assertEqual(duplicate['total_rows'], 1)
        self.assertEqual(duplicate['run_total_rows'], 1)

    def test_old_per_run_files_are_merged_once_into_all_time_output(self):
        os.makedirs(self.batch_dir, exist_ok=True)
        legacy_path = os.path.join(
            self.batch_dir, 'IC_OC_Output_20260914_090000_abcdef.xlsx')
        legacy = Workbook()
        sheet = legacy.active
        sheet.title = 'IC_OC_Output'
        sheet.append(['SĐT', 'Ghi chú input', 'Kết quả', 'Thời gian', 'User chạy'])
        sheet.append([
            '84933333333', 'old-ticket', 'Thành công | Mở OC',
            '14/09/2026 09:00:00', 'old-user',
        ])
        legacy.save(legacy_path)
        legacy.close()

        app_module._ensure_icoc_batch_files()
        app_module._ensure_icoc_batch_files()

        total = load_workbook(self.output_file, data_only=True)
        sheet = total['IC_OC_Output']
        self.assertEqual(sheet.max_row, 2)
        self.assertEqual(sheet['A2'].value, '84933333333')
        self.assertEqual(sheet['B2'].value, 'Thành công | Mở OC')
        self.assertEqual(sheet['D2'].value, 'old-user')
        total.close()

    def test_corrupt_startup_workbooks_are_quarantined_and_recreated(self):
        os.makedirs(self.batch_dir, exist_ok=True)
        with open(self.input_file, 'wb') as file_handle:
            file_handle.write(b'not-an-xlsx-input')
        with open(self.output_file, 'wb') as file_handle:
            file_handle.write(b'<html>not-an-xlsx-output</html>')

        app_module._ensure_icoc_batch_files()

        input_book = load_workbook(self.input_file, data_only=True)
        self.assertIn('IC_OC_Input', input_book.sheetnames)
        self.assertIn('Huong_dan', input_book.sheetnames)
        input_book.close()

        output_book = load_workbook(self.output_file, data_only=True)
        self.assertIn('IC_OC_Output', output_book.sheetnames)
        self.assertEqual(
            [cell.value for cell in output_book['IC_OC_Output'][1]],
            list(app_module.ICOC_BATCH_OUTPUT_HEADERS),
        )
        output_book.close()

        quarantined = os.listdir(self.batch_dir)
        self.assertTrue(any(
            name.startswith('IC_OC_Input.xlsx.corrupt-')
            for name in quarantined))
        self.assertTrue(any(
            name.startswith('IC_OC_Output.xlsx.corrupt-')
            for name in quarantined))


if __name__ == '__main__':
    unittest.main()
