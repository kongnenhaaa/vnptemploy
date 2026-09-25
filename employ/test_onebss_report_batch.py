import base64
import io
import json
import os
import tempfile
import time
import unittest
from datetime import date
from unittest import mock

from openpyxl import Workbook, load_workbook

import onebss_report_batch as report_batch


def _jwt(claims):
    encoded = base64.urlsafe_b64encode(
        json.dumps(claims).encode("utf-8")
    ).decode("ascii").rstrip("=")
    return f"header.{encoded}.signature"


def _report_workbook(rows):
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["BÁO CÁO ĐƠN HÀNG"])
    sheet.append(["DONVI", "NGUONHANG", "MADONHANG", "SANPHAM"])
    for row in rows:
        sheet.append(row)
    stream = io.BytesIO()
    workbook.save(stream)
    workbook.close()
    return stream.getvalue()


class _Response:
    status_code = 200
    headers = {
        "Content-Type": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    }

    def __init__(self, content):
        self.content = content
        self.text = ""

    def json(self):
        raise ValueError


class _HttpSession:
    def __init__(self, response):
        self.response = response
        self.calls = []

    def post(self, url, **kwargs):
        self.calls.append((url, kwargs))
        return self.response


class _JsonResponse:
    status_code = 200
    headers = {"Content-Type": "application/json"}
    text = ""
    content = b""

    def __init__(self, payload):
        self.payload = payload

    def json(self):
        return self.payload


class OneBssReportBatchTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.session_file = os.path.join(self.temp.name, "onebss_session.json")
        token = _jwt({
            "exp": int(time.time()) + 3600,
            "user_name": "tester.hcm",
            "id_tinhthanh": "28",
        })
        with open(self.session_file, "w", encoding="utf-8") as handle:
            json.dump({"sessions": {"active": {"token": {"access_token": token}}}}, handle)

    def tearDown(self):
        self.temp.cleanup()

    def test_download_uses_confirmed_report_filters(self):
        fake = _HttpSession(_Response(_report_workbook([])))
        content = report_batch.download_report_day(
            self.session_file,
            date(2026, 4, 3),
            http_session=fake,
        )
        self.assertTrue(content.startswith(b"PK"))
        _, request = fake.calls[0]
        params = request["json"]["params"]
        self.assertEqual(request["json"]["baocao_id"], 47363)
        self.assertEqual(params["P_PHANVUNG_ID"], "28")
        self.assertEqual(params["P_UNGDUNG"], "19")
        self.assertEqual(params["P_DICHVUVT_ID"], "2")
        self.assertEqual(params["P_LOAITB_ID"], "21")
        self.assertEqual(params["P_HINHTHUC_TT"], "0")
        self.assertEqual(params["P_TUNGAY"], "03/04/2026")
        self.assertEqual(params["P_DENNGAY"], "03/04/2026")
        self.assertEqual(params["username"], "tester.hcm")

    def test_extract_and_merge_deduplicates_rows(self):
        content = _report_workbook([
            ["HCM", "VNPT Employee", "DH-1", "SP-1"],
            ["HCM", "VNPT Employee", "DH-2", "SP-2"],
        ])
        headers, rows = report_batch.extract_report_table(content)
        self.assertEqual(headers, ["DONVI", "NGUONHANG", "MADONHANG", "SANPHAM"])
        output = os.path.join(self.temp.name, "merged.xlsx")
        self.assertEqual(report_batch.merge_report_rows(output, headers, rows), 2)
        self.assertEqual(report_batch.merge_report_rows(output, headers, rows), 0)

        workbook = load_workbook(output, data_only=True)
        sheet = workbook["Du_lieu"]
        self.assertEqual(sheet.max_row, 3)
        self.assertEqual(sheet["C2"].value, "DH-1")
        self.assertFalse(sheet["A2"].alignment.wrap_text)
        workbook.close()

    def test_grid_api_skips_type_row_and_returns_business_rows(self):
        fake = _HttpSession(_JsonResponse({
            "error_code": "BSS-00000000",
            "data": [
                {"DONVI": "String", "NGUONHANG": "String", "MADONHANG": "String"},
                {"DONVI": "HCM", "NGUONHANG": "VNPT Employee", "MADONHANG": "DH-1"},
            ],
        }))
        headers, rows = report_batch.fetch_report_day_table(
            self.session_file,
            date(2026, 8, 1),
            http_session=fake,
        )
        self.assertEqual(headers, ["DONVI", "NGUONHANG", "MADONHANG"])
        self.assertEqual(rows, [["HCM", "VNPT Employee", "DH-1"]])
        self.assertIn("run_v7", fake.calls[0][0])

    def test_manager_runs_each_day_and_removes_checkpoint(self):
        output_dir = os.path.join(self.temp.name, "out")
        manager = report_batch.OneBssReportBatchManager(self.session_file, output_dir)

        def fake_download(_session_file, report_day, **_kwargs):
            return _report_workbook([[
                "HCM", "VNPT Employee", f"DH-{report_day.day}", "SP"
            ]])

        def fake_table(_session_file, report_day, **_kwargs):
            headers, rows = report_batch.extract_report_table(
                fake_download(_session_file, report_day))
            return headers, rows

        with mock.patch.object(report_batch, "fetch_report_day_table", side_effect=fake_table):
            manager.start(date(2026, 4, 1), date(2026, 4, 3))
            manager._thread.join(timeout=10)

        state = manager.status()
        self.assertEqual(state["status"], "completed")
        self.assertEqual(state["completed_days"], 3)
        self.assertEqual(state["rows"], 3)
        self.assertTrue(os.path.isfile(state["output_path"]))
        checkpoints = [name for name in os.listdir(output_dir) if name.endswith("checkpoint.json")]
        self.assertEqual(checkpoints, [])

        restarted = report_batch.OneBssReportBatchManager(self.session_file, output_dir)
        restored = restarted.status()
        self.assertTrue(restored["download_ready"])
        self.assertEqual(restored["output_path"], state["output_path"])


if __name__ == "__main__":
    unittest.main()

